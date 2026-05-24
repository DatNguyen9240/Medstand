import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# ═══ STYLING UTILITIES FOR WORD ═══

def set_cell_shading(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_borders(cell, hex_color="808080", sz="4", val="single"):
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(tcBorders)

def apply_text_formatting(paragraph, text, font_name="Arial", font_size=Pt(10), color_rgb=None):
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            inner = part[2:-2]
            run = paragraph.add_run(inner)
            run.bold = True
            run.font.name = font_name
            run.font.size = font_size
            if color_rgb:
                run.font.color.rgb = color_rgb
        elif part.startswith('`') and part.endswith('`'):
            inner = part[1:-1]
            run = paragraph.add_run(inner)
            run.font.bold = True
            run.font.name = "Courier New"
            run.font.size = font_size - Pt(0.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
        else:
            run = paragraph.add_run(part)
            run.font.name = font_name
            run.font.size = font_size
            if color_rgb:
                run.font.color.rgb = color_rgb

def render_word_table(doc, rows_data, font_family, color_text):
    cleaned_rows = []
    for r in rows_data:
        if re.search(r'^[|\s:-]+$', r):
            continue
        cleaned_rows.append(r)
        
    if not cleaned_rows:
        return
        
    parsed_table = []
    for r in cleaned_rows:
        cells = [c.strip() for c in r.split('|')[1:-1]]
        parsed_table.append(cells)
        
    if not parsed_table:
        return
        
    num_rows = len(parsed_table)
    num_cols = len(parsed_table[0])
    
    w_table = doc.add_table(rows=num_rows, cols=num_cols)
    w_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_table.autofit = True
    
    for r_idx, r_data in enumerate(parsed_table):
        row = w_table.rows[r_idx]
        is_header = (r_idx == 0)
        shading_color = "F2F2F2" if is_header else "FFFFFF"
        
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            set_cell_shading(cell, shading_color)
            set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
            
            border_color = "808080" if is_header else "D3D3D3"
            set_cell_borders(cell, hex_color=border_color, sz="4")
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            if len(val) <= 4 or val.isdigit() or (is_header and val.upper() in ["STT", "MÃ SP", "ĐẠT", "KHÔNG ĐẠT", "TÀI KHOẢN QL", "TÀI KHOẢN TDV", "MÃ TÀI KHOẢN"]):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            run = p.add_run()
            run.font.name = font_family
            run.font.size = Pt(9)
            
            if is_header:
                run.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.text = val
            else:
                apply_text_formatting(p, val, font_family, Pt(9), color_text)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ═══ JS DATA PARSER ═══

def parse_js_data(text):
    users = {}
    blocks = re.findall(r'"([^"]+)":\s*{(.*?)\s*prompts:\s*\[(.*?)\s*\]\s*}', text, re.DOTALL)
    for username, meta_text, prompts_text in blocks:
        meta = {}
        for k, v in re.findall(r'(\w+):\s*"([^"]+)"', meta_text):
            meta[k] = v
            
        prompts = []
        for prompt_match in re.findall(r'\{\s*stt:\s*"([^"]+)",\s*name:\s*"([^"]+)",\s*text:\s*"([^"]+)"\s*\}', prompts_text):
            prompts.append({
                'stt': prompt_match[0],
                'name': prompt_match[1],
                'text': prompt_match[2]
            })
        meta['prompts'] = prompts
        users[username] = meta
    return users

# ═══ MAIN Compile ═══

def run_restructuring():
    md_path = r"c:\Users\Legion\Desktop\AI Nhà Thuốc\Medstand\HDSD_Medstand_AI_Business.md"
    docx_path = r"c:\Users\Legion\Desktop\AI Nhà Thuốc\Medstand\HDSD_Medstand_AI_Business.docx"
    docx_alt_path = r"c:\Users\Legion\Desktop\AI Nhà Thuốc\Medstand\HDSD_Medstand_AI_Business_Updated.docx"
    
    print("Reading MD manual...")
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()
        
    print("Compiling to Word DOCX...")
    try:
        compile_to_docx(md_content, docx_path)
        print("Primary DOCX compiled successfully!")
    except PermissionError:
        print("\n[Warning] Primary path locked! Saving to alternative path...")
        compile_to_docx(md_content, docx_alt_path)
        print("Fallback DOCX compiled successfully!")


def render_callout_box(doc, lines_data, font_family):
    w_table = doc.add_table(rows=1, cols=1)
    w_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_table.autofit = True
    
    cell = w_table.rows[0].cells[0]
    set_cell_shading(cell, "FAFAFA")  # Soft light gray shading
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_borders(cell, hex_color="D0D3D4", sz="6", val="dashed")  # Nice dashed border
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    for i, line in enumerate(lines_data):
        if i > 0:
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
        apply_text_formatting(p, line, font_family, Pt(9.5), RGBColor(100, 100, 100))
        
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)


def compile_to_docx(md_content, out_docx_path):
    doc = docx.Document()
    
    # Page setup (margins)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        
    FONT_FAMILY = "Arial"
    COLOR_PRIMARY = RGBColor(0, 0, 0)
    COLOR_SECONDARY = RGBColor(0, 0, 0)
    COLOR_TEXT = RGBColor(0, 0, 0)
    
    lines = md_content.split('\n')
    
    in_table = False
    table_rows = []
    
    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        
        # Table parsing
        if line.startswith('|'):
            in_table = True
            table_rows.append(line)
            idx += 1
            continue
        elif in_table:
            in_table = False
            render_word_table(doc, table_rows, FONT_FAMILY, COLOR_TEXT)
            table_rows = []
            
        # Blockquote (Callout / Image Slot)
        if line.startswith('>'):
            blockquote_lines = []
            while idx < len(lines) and lines[idx].strip().startswith('>'):
                clean_line = lines[idx].strip()[1:].strip()
                blockquote_lines.append(clean_line)
                idx += 1
            render_callout_box(doc, blockquote_lines, FONT_FAMILY)
            continue
            
        if not line or line.startswith('---'):
            idx += 1
            continue
            
        # Heading 1
        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = FONT_FAMILY
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            
        # Heading 2
        elif line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(15)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = FONT_FAMILY
            run.font.size = Pt(13.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_SECONDARY
            
        # Heading 3
        elif line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = FONT_FAMILY
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLOR_TEXT
            
        # Heading 4 (sub-sections)
        elif line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = FONT_FAMILY
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = COLOR_TEXT
            
        # Bullet list item
        elif line.startswith('- ') or line.startswith('* '):
            text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3.5)
            apply_text_formatting(p, text, FONT_FAMILY, Pt(10), COLOR_TEXT)
            
        # Numbered list item
        elif re.match(r'^\d+\.\s', line):
            match = re.match(r'^(\d+\.)\s(.*)', line)
            num = match.group(1)
            text = match.group(2).strip()
            
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3.5)
            
            run_num = p.add_run(num + " ")
            run_num.font.name = FONT_FAMILY
            run_num.font.size = Pt(10)
            run_num.font.bold = True
            run_num.font.color.rgb = COLOR_PRIMARY
            
            apply_text_formatting(p, text, FONT_FAMILY, Pt(10), COLOR_TEXT)
            
        # Standard paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(5)
            apply_text_formatting(p, line, FONT_FAMILY, Pt(10), COLOR_TEXT)
            
        idx += 1
        
    if in_table and table_rows:
        render_word_table(doc, table_rows, FONT_FAMILY, COLOR_TEXT)
        
    print("Saving beautiful DOCX...")
    doc.save(out_docx_path)
    print("DOCX saved successfully!")

if __name__ == "__main__":
    run_restructuring()
